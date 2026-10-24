#!/usr/bin/env python3
"""
Amex-export -> taggade rader, och perioddata -> månadsrapport.

Körs med openpyxl via uv:
  uv run --no-project --with openpyxl python amex.py tag <export.xlsx|katalog> --rules R.json --known K.json
  uv run --no-project --with openpyxl python amex.py report --tx T.json --period 2026-10 [--settlement S.json]

tag
  <export>   En Amex-export (.xlsx), eller en katalog: då väljs den nyaste .xlsx vars
             rubrikrad ser ut som en Amex-export. Filnamnet spelar ingen roll.
  --rules    Regler-fliken som JSON: lista av rader
             [Mönster, Kortmedlem, Tagg, Kategori, Företag, Säkerhet, Historik].
  --known    Referenser som redan finns i Transaktioner (JSON-lista). Får saknas.
  Ut: JSON {periods, rows, duplicates, counts}. Varje rad har ask=true när den ska frågas.

report
  --tx       Transaktioner som JSON: lista av dicts med period, date, desc, card, amount,
             tag, category, foretag.
  --period   Perioden rapporten gäller.
  --settlement  Avräkningsraden för perioden som JSON (valfri), skrivs in i rapporten.
  Ut: markdown med siffrorna; iakttagelserna skriver Claude själv.
"""

import argparse
import collections
import datetime
import json
import re
import sys
from pathlib import Path

PERIOD_END_DAY = 24
REQUIRED = {"Datum", "Beskrivning", "Kortmedlem", "Belopp"}


def norm(text):
    return re.sub(r"\s+", " ", str(text or "").replace("\xa0", " ")).strip().upper()


def to_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    m = re.match(r"(\d\d)/(\d\d)/(\d{4})", str(value or ""))
    return datetime.date(int(m[3]), int(m[1]), int(m[2])) if m else None


def period_of(day):
    """Perioden heter efter månaden den slutar i: 25 sep - 24 okt -> 2026-10."""
    if day.day > PERIOD_END_DAY:
        day = (day.replace(day=1) + datetime.timedelta(days=32)).replace(day=1)
    return f"{day.year}-{day.month:02d}"


def find_header(ws):
    for r in range(1, 15):
        names = {str(ws.cell(r, c).value).strip(): c for c in range(1, 30) if ws.cell(r, c).value}
        if REQUIRED <= names.keys():
            return r, names
    return None, None


def pick_export(path):
    import openpyxl
    path = Path(path).expanduser()
    if path.is_file():
        return path
    for f in sorted(path.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True):
        try:
            ws = openpyxl.load_workbook(f, read_only=False, data_only=True).active
        except Exception:
            continue
        if find_header(ws)[0]:
            return f
    sys.exit(f"Ingen Amex-export hittades i {path}")


def parse(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).active
    hr, col = find_header(ws)
    if not hr:
        sys.exit(f"{path} ser inte ut som en Amex-export")
    ref_col = col.get("Referens")
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        day = to_date(ws.cell(r, col["Datum"]).value)
        amount = ws.cell(r, col["Belopp"]).value
        if not day or amount in (None, ""):
            continue
        if norm(ws.cell(r, col["Beskrivning"]).value).startswith("BETALNING MOTTAGEN"):
            continue
        if isinstance(amount, str):
            amount = float(amount.replace(" ", "").replace(",", "."))
        ref = ws.cell(r, ref_col).value if ref_col else None
        desc = re.sub(r"\s+", " ", str(ws.cell(r, col["Beskrivning"]).value).replace("\xa0", " ")).strip()
        card = norm(ws.cell(r, col["Kortmedlem"]).value)
        rows.append({
            "period": period_of(day),
            "date": day.isoformat(),
            "desc": desc,
            "card": card,
            "amount": round(float(amount), 2),
            "ref": str(ref).strip() if ref else f"{day.isoformat()}|{desc}|{card}|{amount}",
        })
    return rows


def match(rules, desc, card):
    """Längsta matchande mönster vinner; en kortmedlemsregel går före en allmän."""
    d = norm(desc)
    best = None
    for pattern, rule_card, tag, category, foretag, certainty, *history in rules:
        if norm(pattern) in d and (not rule_card or norm(rule_card) == card):
            rank = (len(norm(pattern)), bool(rule_card))
            if not best or rank > best[0]:
                best = (rank, {"pattern": pattern, "tag": tag, "category": category,
                               "foretag": foretag, "certainty": certainty,
                               "history": history[0] if history else ""})
    return best[1] if best else None


def cmd_tag(args):
    export = pick_export(args.export)
    rules = json.loads(Path(args.rules).read_text())
    known = set(json.loads(Path(args.known).read_text())) if args.known else set()
    rows, dups = [], []
    for row in parse(export):
        if row["ref"] in known:
            dups.append(row)
            continue
        known.add(row["ref"])
        hit = match(rules, row["desc"], row["card"])
        if hit:
            row.update(tag=hit["tag"], category=hit["category"], foretag=hit["foretag"],
                       source="regel", rule=hit["pattern"], history=hit["history"],
                       ask=hit["certainty"] != "säker")
        else:
            row.update(tag="", category="", foretag="", source="", rule="", history="", ask=True)
        rows.append(row)
    out = {
        "export": str(export),
        "periods": dict(collections.Counter(r["period"] for r in rows)),
        "counts": {"rows": len(rows), "ask": sum(r["ask"] for r in rows), "duplicates": len(dups)},
        "rows": rows,
        "duplicates": dups,
    }
    json.dump(out, sys.stdout, ensure_ascii=False, indent=1)


def merchant(desc):
    d = re.sub(r"^(PAYPAL|ZETTLE|SUMUP|CKO|PAYEX|IZ|SQ|CPAYMENT|GOOGLE)\s*\*\s*", "", norm(desc))
    words = re.findall(r"[A-ZÅÄÖÉ&]{2,}", d)
    return " ".join(words[:2]) if words else d


def kr(x):
    return f"{x:,.0f}".replace(",", " ") + " kr"


def previous_periods(period, n):
    y, m = map(int, period.split("-"))
    out = []
    for _ in range(n):
        m -= 1
        if m == 0:
            y, m = y - 1, 12
        out.append(f"{y}-{m:02d}")
    return out


def cmd_report(args):
    tx = json.loads(Path(args.tx).read_text())
    p = args.period
    rows = [t for t in tx if t["period"] == p]
    prev = previous_periods(p, 3)
    by_cat = collections.defaultdict(float)
    by_cat_prev = collections.defaultdict(float)
    for t in rows:
        by_cat[t["category"] or "Övrigt"] += t["amount"]
    prev_present = sorted({t["period"] for t in tx if t["period"] in prev})
    for t in tx:
        if t["period"] in prev_present:
            by_cat_prev[t["category"] or "Övrigt"] += t["amount"] / len(prev_present)

    people = collections.defaultdict(lambda: collections.defaultdict(float))
    for t in rows:
        people[t["card"]][t["tag"] or "GEMENSAM"] += t["amount"]

    lines = [f"## Översikt", "",
             "| Kortmedlem | Gemensamt | Privat | Totalt |", "|---|---:|---:|---:|"]
    for card in sorted(people):
        g, pr = people[card]["GEMENSAM"], people[card]["PRIVAT"]
        lines.append(f"| {card.title()} | {kr(g)} | {kr(pr)} | {kr(g + pr)} |")
    total = sum(t["amount"] for t in rows)
    common = sum(t["amount"] for t in rows if t["tag"] != "PRIVAT")
    lines += [f"| **Alla** | **{kr(common)}** | **{kr(total - common)}** | **{kr(total)}** |", "",
              "Taggarna som de står i Transaktioner. Lunchavdraget flyttas först i avräkningen.", ""]

    if args.settlement:
        s = json.loads(Path(args.settlement).read_text())
        lines += ["## Avräkning", ""] + [f"- {k}: {kr(v) if isinstance(v, (int, float)) else v}" for k, v in s.items()] + [""]

    avg_label = f"snitt {prev_present[0]}–{prev_present[-1]}" if prev_present else "snitt"
    lines += ["## Kategorier", "", f"| Kategori | Gemensamt | Privat | Totalt | {avg_label} | Δ |", "|---|---:|---:|---:|---:|---:|"]
    cat_split = collections.defaultdict(lambda: [0.0, 0.0])
    for t in rows:
        cat_split[t["category"] or "Övrigt"][t["tag"] == "PRIVAT"] += t["amount"]
    for cat, v in sorted(by_cat.items(), key=lambda kv: -kv[1]):
        avg = by_cat_prev.get(cat, 0.0)
        delta = f"{(v - avg) / avg * 100:+.0f} %" if avg > 0 and prev_present else "–"
        lines.append(f"| {cat} | {kr(cat_split[cat][0])} | {kr(cat_split[cat][1])} | {kr(v)} | {kr(avg) if prev_present else '–'} | {delta} |")
    lines.append("")

    merchants = collections.defaultdict(lambda: [0, 0.0])
    for t in rows:
        key = merchant(t["desc"])
        merchants[key][0] += 1
        merchants[key][1] += t["amount"]
    lines += ["## Största handlare", "", "| Handlare | Antal | Summa |", "|---|---:|---:|"]
    for k, (n, s) in sorted(merchants.items(), key=lambda kv: -kv[1][1])[:10]:
        lines.append(f"| {k.title()} | {n} | {kr(s)} |")
    lines.append("")

    subs = [t for t in rows if t["category"] == "Prenumerationer" or t.get("foretag")]
    if subs:
        seen_before = {merchant(t["desc"]).split()[0] for t in tx if t["period"] in prev}
        lines += ["## Prenumerationer och företagskandidater", "", "| Tjänst | Belopp | Tagg | Not |", "|---|---:|---|---|"]
        for t in sorted(subs, key=lambda t: -t["amount"]):
            key = merchant(t["desc"]).split()[0]
            notes = []
            if t.get("foretag"):
                notes.append("företagskandidat")
            if key not in seen_before:
                notes.append("ny")
            lines.append(f"| {t['desc'][:40]} | {kr(t['amount'])} | {t['tag'].lower()} | {', '.join(notes)} |")
        fsum = sum(t["amount"] for t in rows if t.get("foretag"))
        fall = sum(t["amount"] for t in tx if t.get("foretag") and t["period"] <= p)
        lines += ["", f"Företagskandidater: {kr(fsum)} denna period, {kr(fall)} totalt sedan start.", ""]

    print("\n".join(lines))


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("tag")
    t.add_argument("export", nargs="?", default="~/Downloads")
    t.add_argument("--rules", required=True)
    t.add_argument("--known")
    r = sub.add_parser("report")
    r.add_argument("--tx", required=True)
    r.add_argument("--period", required=True)
    r.add_argument("--settlement")
    args = ap.parse_args()
    {"tag": cmd_tag, "report": cmd_report}[args.cmd](args)


if __name__ == "__main__":
    main()
